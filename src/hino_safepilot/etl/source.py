"""Immutable source validation and bounded-memory workbook iteration."""

from __future__ import annotations

import hashlib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Iterator

from openpyxl import load_workbook

from .constants import (
    EXPECTED_HEADERS,
    EXPECTED_MAPPING_DIMENSION,
    EXPECTED_SOURCE_COLUMNS,
    EXPECTED_SOURCE_DIMENSION,
    EXPECTED_SOURCE_ROWS,
    MAPPING_SHEET,
    SOURCE_SHEET,
)


class SourceValidationError(RuntimeError):
    """Raised when immutable source expectations or schema contracts fail."""


@dataclass(frozen=True)
class SourceMetadata:
    path: Path
    file_name: str
    sha256: str
    size_bytes: int
    mtime_ns: int
    sheet: str
    source_rows: int
    source_columns: int
    source_dimension: str
    mapping_dimension: str
    headers: tuple[str, ...]


def sha256_file(path: Path, chunk_size: int = 1024 * 1024) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        while chunk := stream.read(chunk_size):
            digest.update(chunk)
    return digest.hexdigest().upper()


def expected_hash_from_manifest(manifest_path: Path, source_relative_path: str) -> str:
    target = source_relative_path.replace("\\", "/")
    for line in manifest_path.read_text(encoding="utf-8").splitlines():
        if target in line.replace("\\", "/"):
            matches = re.findall(r"`([A-Fa-f0-9]{64})`", line)
            if matches:
                return matches[-1].upper()
    raise SourceValidationError(
        f"Source SHA-256 for {source_relative_path!r} was not found in {manifest_path}."
    )


def validate_source(
    path: Path,
    *,
    expected_sha256: str,
    expected_rows: int = EXPECTED_SOURCE_ROWS,
    expected_columns: int = EXPECTED_SOURCE_COLUMNS,
    expected_headers: tuple[str, ...] = EXPECTED_HEADERS,
    expected_source_dimension: str = EXPECTED_SOURCE_DIMENSION,
    expected_mapping_dimension: str = EXPECTED_MAPPING_DIMENSION,
) -> SourceMetadata:
    if not path.is_file():
        raise SourceValidationError(f"Source workbook not found: {path}")
    stat_before = path.stat()
    actual_hash = sha256_file(path)
    if actual_hash != expected_sha256.upper():
        raise SourceValidationError(
            f"Source SHA-256 mismatch: expected {expected_sha256.upper()}, got {actual_hash}."
        )

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        missing_sheets = {SOURCE_SHEET, MAPPING_SHEET}.difference(workbook.sheetnames)
        if missing_sheets:
            raise SourceValidationError(
                f"Required sheet(s) missing: {sorted(missing_sheets)}; found {workbook.sheetnames}."
            )
        sheet = workbook[SOURCE_SHEET]
        mapping = workbook[MAPPING_SHEET]
        headers = tuple(cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1)))
        duplicate_headers = sorted({h for h in headers if headers.count(h) > 1})
        if duplicate_headers:
            raise SourceValidationError(f"Duplicate source columns: {duplicate_headers}")
        if headers != expected_headers:
            missing = sorted(set(expected_headers).difference(headers))
            extra = sorted(set(headers).difference(expected_headers))
            raise SourceValidationError(
                f"Source schema drift. missing={missing}, extra={extra}, ordered_headers_changed={not missing and not extra}."
            )
        actual_rows = sheet.max_row - 1
        if actual_rows != expected_rows or sheet.max_column != expected_columns:
            raise SourceValidationError(
                f"Source shape drift: expected {expected_rows}x{expected_columns}, "
                f"got {actual_rows}x{sheet.max_column}."
            )
        source_dimension = sheet.calculate_dimension()
        mapping_dimension = mapping.calculate_dimension()
        if source_dimension != expected_source_dimension:
            raise SourceValidationError(
                f"Source dimension drift: expected {expected_source_dimension}, got {source_dimension}."
            )
        if mapping_dimension != expected_mapping_dimension:
            raise SourceValidationError(
                f"Mapping dimension drift: expected {expected_mapping_dimension}, got {mapping_dimension}."
            )
    finally:
        workbook.close()

    stat_after = path.stat()
    if stat_before.st_mtime_ns != stat_after.st_mtime_ns or stat_before.st_size != stat_after.st_size:
        raise SourceValidationError("Source file metadata changed during validation.")
    return SourceMetadata(
        path=path,
        file_name=path.name,
        sha256=actual_hash,
        size_bytes=stat_after.st_size,
        mtime_ns=stat_after.st_mtime_ns,
        sheet=SOURCE_SHEET,
        source_rows=actual_rows,
        source_columns=sheet.max_column,
        source_dimension=source_dimension,
        mapping_dimension=mapping_dimension,
        headers=headers,
    )


def iter_source_rows(path: Path) -> Iterator[tuple[int, tuple[Any, ...]]]:
    """Yield one source row at a time; row numbers use Excel's 1-based coordinates."""

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook[SOURCE_SHEET]
        rows = sheet.iter_rows(values_only=True)
        headers = tuple(next(rows))
        if headers != EXPECTED_HEADERS:
            raise SourceValidationError("Source headers changed between validation and streaming.")
        width = len(EXPECTED_HEADERS)
        for source_row_number, row in enumerate(rows, start=2):
            if len(row) < width:
                row = tuple(row) + (None,) * (width - len(row))
            elif len(row) > width:
                raise SourceValidationError(
                    f"Row {source_row_number} exceeds expected width {width}."
                )
            yield source_row_number, tuple(row)
    finally:
        workbook.close()

