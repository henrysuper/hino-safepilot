from __future__ import annotations

from copy import deepcopy
from pathlib import Path

import pytest
from openpyxl import Workbook

from hino_safepilot.etl.constants import EXPECTED_HEADERS


def _base_row(**overrides: object) -> dict[str, object]:
    row: dict[str, object] = {header: "" for header in EXPECTED_HEADERS}
    row.update(
        {
            "Type": 300,
            "journeyCode": "250101000001",
            "time": "2025-01-01 00:00:00",
            "carStatus": "1",
            "enabledCode": "SYNTH00001",
            "carNum": "TEST-0001",
            "gps.longitude": "0",
            "gps.latitude": "0",
            "gps.speed": "50",
            "gps.azimuth": "90",
            "gps.speedLimit": "60",
            "can.canStatus": 0,
            "can.totalMileage": "1000.0",
            "can.canSpeed": "49",
            "line.fuelLevel": "80",
            "line.battery": "24000",
            "can.engine.engineCoolantTemp": "70",
            "can.engine.totalFuelUsed": "100.0",
            "can.engine.engineTotalTime": "50.0",
            "can.engine.rpm": "1200",
            "can.engine.engineLoad": "30",
            "fenceId": "",
        }
    )
    row.update(overrides)
    return row


def synthetic_rows() -> list[dict[str, object]]:
    normal = _base_row()
    missing_gps_duplicate = _base_row(
        **{"gps.longitude": None, "gps.latitude": None, "gps.speed": None}
    )
    malformed_numeric = _base_row(
        time="2025-01-01 00:01:00", **{"gps.speed": "not_numeric"}
    )
    same_journey_other_vehicle = _base_row(
        enabledCode="SYNTH00002", carNum="TEST-0002"
    )
    all_slots = _base_row(
        time="2025-01-01 00:04:00",
        **{
            "event[0].type": "6",
            "event[0].startTime": "2025-01-01 00:04:00",
            "event[0].info.speedDifference": "10",
            "event[1].type": "7",
            "event[1].startTime": "2025-01-01 00:04:00",
            "event[1].info.speedDifference": "12",
            "event[2].type": "8",
            "event[2].startTime": "2025-01-01 00:04:00",
            "event[2].endTime": "2025-01-01 00:05:00",
            "event[2].info.duration": 60,
            "event[2].info.speedLimit": "50",
            "event[2].info.speed": "70",
        },
    )
    repeated_event = _base_row(
        time="2025-01-01 00:04:30",
        **{
            "event[0].type": "8",
            "event[0].startTime": "2025-01-01 00:04:00",
            "event[0].endTime": "2025-01-01 00:05:00",
            "event[0].info.duration": 60,
            "event[0].info.speedLimit": "50",
            "event[0].info.speed": "70",
        },
    )
    conflicting_event = _base_row(
        time="2025-01-01 00:05:00",
        **{
            "event[0].type": "8",
            "event[0].startTime": "2025-01-01 00:04:00",
            "event[0].endTime": "2025-01-01 00:06:00",
            "event[0].info.duration": 120,
            "event[0].info.speedLimit": "50",
            "event[0].info.speed": "80",
        },
    )
    out_of_range = _base_row(
        time="2025-01-01 00:06:00",
        **{"gps.speed": "477", "gps.azimuth": "360"},
    )
    dtc = _base_row(
        time="2025-01-01 00:07:00",
        **{
            "event[0].type": "9",
            "event[0].startTime": "2025-01-01 00:07:00",
            "event[0].info.dtcCount": "1",
            "event[0].info.dtcCodes[0]": "PTEST0",
        },
    )
    malformed_timestamp = _base_row(time="not-a-time")
    identical_a = _base_row(
        journeyCode="250101000002", time="2025-01-01 01:00:00"
    )
    identical_b = deepcopy(identical_a)
    return [
        normal,
        missing_gps_duplicate,
        malformed_numeric,
        same_journey_other_vehicle,
        all_slots,
        repeated_event,
        conflicting_event,
        out_of_range,
        dtc,
        malformed_timestamp,
        identical_a,
        identical_b,
    ]


def write_synthetic_workbook(path: Path, rows: list[dict[str, object]] | None = None) -> int:
    rows = synthetic_rows() if rows is None else rows
    workbook = Workbook()
    output = workbook.active
    output.title = "output"
    output.append(list(EXPECTED_HEADERS))
    for row in rows:
        output.append([row.get(header) for header in EXPECTED_HEADERS])
    mapping = workbook.create_sheet("欄位對照表")
    mapping.cell(row=1, column=2, value="synthetic mapping start")
    mapping.cell(row=82, column=8, value="synthetic mapping end")
    workbook.save(path)
    workbook.close()
    return len(rows)


@pytest.fixture
def synthetic_workbook(tmp_path: Path) -> tuple[Path, int]:
    path = tmp_path / "synthetic_telemetry.xlsx"
    count = write_synthetic_workbook(path)
    return path, count

